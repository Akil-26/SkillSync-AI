"""
Stage 1: Offline preprocessing.

Run this ONCE, offline, with network access allowed (to download the
embedding model the first time). Produces all artifacts that rank.py needs
to run within the 5-minute / no-network / CPU-only constraint:

  - artifacts/models/all-MiniLM-L6-v2/      (cached embedding model)
  - artifacts/embeddings/candidate_embeddings.npy + candidate_ids.json
  - artifacts/embeddings/jd_embedding.npy
  - artifacts/features/feature_table.json   (structured features per candidate)

Usage:
    python -m src.precompute
"""

from __future__ import annotations

import json

from . import config, embedding, features
from .utils import get_logger, read_jsonl_gz, safe_get, timer, write_json

logger = get_logger("precompute", log_file=config.LOGS_DIR / "precompute.log")


def load_candidates() -> list[dict]:
    if not config.CANDIDATES_PATH.exists():
        raise FileNotFoundError(
            f"Candidates file not found at {config.CANDIDATES_PATH}. "
            f"Place candidates.jsonl.gz in the data/ directory."
        )
    candidates = read_jsonl_gz(config.CANDIDATES_PATH)
    logger.info(f"Loaded {len(candidates)} candidates from {config.CANDIDATES_PATH}")
    return candidates


def parse_job_description() -> dict:
    """Parse job_description.docx into the structured dict shape defined by
    config.JD_SECTIONS.

    ASSUMPTION: the docx uses simple heading-based sections (e.g. a heading
    'Must Have' followed by bullet points). This is a best-effort parser —
    once the real file is available, verify the heading names match and
    adjust SECTION_HEADING_ALIASES below if needed.
    """
    if not config.JOB_DESCRIPTION_PATH.exists():
        raise FileNotFoundError(
            f"Job description not found at {config.JOB_DESCRIPTION_PATH}."
        )

    # pyrefly: ignore [missing-import]
    import docx  # python-docx

    doc = docx.Document(str(config.JOB_DESCRIPTION_PATH))

    SECTION_HEADING_ALIASES = {
        "must have": "must_have",
        "must-have": "must_have",
        "required": "must_have",
        "requirements": "must_have",
        "nice to have": "nice_to_have",
        "preferred": "nice_to_have",
        "responsibilities": "responsibilities",
        "required skills": "required_skills",
        "skills": "required_skills",
    }

    sections: dict[str, list[str]] = {}
    current_section = "body"
    full_text_parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        full_text_parts.append(text)

        heading_key = text.lower().rstrip(":")
        matched_section = SECTION_HEADING_ALIASES.get(heading_key)
        if matched_section:
            current_section = matched_section
            sections.setdefault(current_section, [])
            continue

        sections.setdefault(current_section, []).append(text)

    full_text = "\n".join(full_text_parts)

    jd = {
        "title": doc.paragraphs[0].text.strip() if doc.paragraphs else "",
        "full_text": full_text,
        config.JD_SECTIONS["must_have"]: sections.get("must_have", []),
        config.JD_SECTIONS["nice_to_have"]: sections.get("nice_to_have", []),
        config.JD_SECTIONS["responsibilities"]: sections.get("responsibilities", []),
        config.JD_SECTIONS["required_skills"]: _flatten_skill_lines(
            sections.get("required_skills", [])
        ),
        config.JD_SECTIONS["min_experience_years"]: _extract_min_experience(full_text),
    }
    logger.info(
        f"Parsed JD: title={jd['title']!r}, "
        f"{len(jd[config.JD_SECTIONS['required_skills']])} required skills detected."
    )
    return jd


def _flatten_skill_lines(lines: list[str]) -> list[str]:
    skills = []
    for line in lines:
        # split on commas / bullets, since skills are often comma-separated
        parts = [p.strip("-•* \t") for p in line.replace(";", ",").split(",")]
        skills.extend(p for p in parts if p)
    return skills


def _extract_min_experience(text: str) -> float | None:
    import re

    match = re.search(r"(\d+)\+?\s*(?:years|yrs)", text, re.IGNORECASE)
    if match:
        return float(match.group(1))
    return None


def build_resume_text(candidate: dict) -> str:
    """Concatenate the fields most relevant for semantic matching into one
    string for embedding. Adjust once real schema is confirmed."""
    schema = config.SCHEMA
    from .utils import as_list, safe_get

    parts = [
        safe_get(candidate, schema["resume_text"], "") or "",
        safe_get(candidate, schema["current_title"], "") or "",
        " ".join(as_list(safe_get(candidate, schema["past_titles"], []))),
        " ".join(as_list(safe_get(candidate, schema["skills"], []))),
    ]
    return " ".join(p for p in parts if p).strip()


def run(allow_model_download: bool = True) -> None:
    logger.info("=== Stage 1: Precompute starting ===")

    with timer("load_candidates", logger):
        candidates = load_candidates()

    with timer("parse_job_description", logger):
        jd = parse_job_description()

    with timer("compute_embeddings", logger):
        resume_texts = [build_resume_text(c) for c in candidates]
        candidate_ids = [safe_get(c, config.SCHEMA["id"]) for c in candidates]

        candidate_embeddings = embedding.encode_texts(
            resume_texts, allow_download=allow_model_download
        )
        jd_embedding = embedding.encode_texts(
            [jd.get("full_text", "")], allow_download=allow_model_download
        )[0]

        embedding.save_embeddings(
            candidate_embeddings,
            candidate_ids,
            config.CANDIDATE_EMBEDDINGS_PATH,
            config.CANDIDATE_IDS_PATH,
        )
        # pyrefly: ignore [missing-import]
        import numpy as np
        np.save(config.JD_EMBEDDING_PATH, jd_embedding)

    with timer("compute_structured_features", logger):
        duplicate_map = features.detect_duplicate_resumes(candidates)
        feature_table = [features.compute_all_features(c, jd) for c in candidates]
        for row in feature_table:
            cid = row["candidate_id"]
            if cid in duplicate_map:
                row["duplicate_of"] = duplicate_map[cid]

        config.FEATURES_DIR.mkdir(parents=True, exist_ok=True)
        write_json(config.FEATURES_DIR / "feature_table.json", feature_table)

    # Cache the parsed JD too, so rank.py doesn't need python-docx at runtime
    write_json(config.FEATURES_DIR / "parsed_jd.json", jd)

    logger.info("=== Stage 1: Precompute complete ===")
    logger.info(f"Artifacts written to {config.ARTIFACTS_DIR}")


if __name__ == "__main__":
    run()